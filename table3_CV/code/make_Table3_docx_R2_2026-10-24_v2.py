"""
Table 3 — build the editable Word table from the R output.
Version 2026-10-24 v2 (CAI row removed) (R2 revision, PLOS ONE PONE-D-26-36808)

Input : table3_CV/output/Table3_CV_wide_R2_2026-10-24.csv  (made by Table3_R2_2026-10-24_v2.R)
Output: table3_CV/output/Table3_R2_2026-10-24.docx

Run after the R script:  python3 make_Table3_docx_R2_2026-10-24.py
Needs: pip install python-docx
No numbers are typed here — every value is read from the CSV and printed to 1 decimal.
"""
import csv
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# ---- TUNING BLOCK ----
BASE = "/Users/gbbulut/Desktop/making figures attempt 14/dogfood-AGE-figures/table3_CV/output/"
IN_CSV = BASE + "Table3_CV_wide_R2_2026-10-24.csv"
OUT_DOCX = BASE + "Table3_R2_2026-10-24.docx"
FONT = "Arial"
BASES = ["Dry matter (µg/g)", "As-fed (µg/g)", "Caloric basis (µg/kcal ME)"]
CAPTION = ("Coefficient of variation (CV, %) of AGE-associated measures in wet (n = 56) and "
           "dry (kibble; n = 44) dog foods, by reporting basis.")
FOOTNOTE = ("CV = standard deviation / mean × 100. Hydrophilic + hydrophobic = sum of the two "
            "fluorescence measures. ME, metabolizable energy.")
NOTE = ""   # internal note to co-authors removed (it was deleted by hand from the submitted table)
# ----------------------


def set_border(cell, top=False, bottom=False):
    """Thin grey top/bottom rules only (booktabs style)."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side, on in (("top", top), ("bottom", bottom), ("left", False), ("right", False)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single" if on else "nil")
        if on:
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "808080")
        borders.append(el)
    tcPr.append(borders)


def write(cell, text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=9):
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.name, r.font.size = FONT, Pt(size)


with open(IN_CSV, encoding="utf-8") as f:
    rows = list(csv.reader(f))
header, body = rows[0], rows[1:]
assert len(header) == 7, "expected Metric + 6 value columns"

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Inches(11), Inches(8.5)
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, m, Inches(0.75))

cap = doc.add_paragraph()
r = cap.add_run("Table 3. "); r.bold = True; r.font.name = FONT; r.font.size = Pt(10)
r = cap.add_run(CAPTION); r.font.name = FONT; r.font.size = Pt(10)

t = doc.add_table(rows=2 + len(body), cols=7)
t.alignment = WD_TABLE_ALIGNMENT.LEFT
t.autofit = False
widths = [Inches(3.0)] + [Inches(1.0)] * 6
for i, w in enumerate(widths):
    t.columns[i].width = w
for row in t.rows:
    for i, w in enumerate(widths):
        row.cells[i].width = w

# header row 1: basis spanning Wet / Dry
write(t.cell(0, 0), "AGE-associated measure", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
set_border(t.cell(0, 0), top=True)
for k, b in enumerate(BASES):
    c = t.cell(0, 1 + 2 * k).merge(t.cell(0, 2 + 2 * k))
    write(c, b, bold=True)
    set_border(c, top=True, bottom=True)
# header row 2
set_border(t.cell(1, 0), bottom=True)
for j in range(6):
    c = t.cell(1, 1 + j)
    write(c, "Wet" if j % 2 == 0 else "Dry (kibble)", bold=True)
    set_border(c, bottom=True)

# body: values straight from the CSV
for i, r_ in enumerate(body):
    optional = "[optional row]" in r_[0]
    last = i == len(body) - 1
    c = t.cell(2 + i, 0)
    write(c, r_[0].replace(" [optional row]", ""), italic=optional, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_border(c, bottom=last)
    for j, v in enumerate(r_[1:]):
        c = t.cell(2 + i, 1 + j)
        write(c, f"{float(v):.1f}", italic=optional)
        set_border(c, bottom=last)

fn = doc.add_paragraph()
r = fn.add_run(FOOTNOTE); r.font.name = FONT; r.font.size = Pt(8)
if NOTE:
    nt = doc.add_paragraph()
    r = nt.add_run(NOTE); r.italic = True; r.font.name = FONT; r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

doc.save(OUT_DOCX)
print("wrote", OUT_DOCX)
